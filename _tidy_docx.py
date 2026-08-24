import shutil, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.docx'
DST = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍_整理版.docx'

shutil.copy2(SRC, DST)
doc = Document(DST)

GREEN = RGBColor(0x22, 0xC5, 0x5E)
DARK  = RGBColor(0x1F, 0x2A, 0x37)
GRAY  = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def set_run_font(run, name='微软雅黑', size=None, color=None, bold=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name); rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name)
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.font.bold = bool(bold)

def style_set_font(style, name='微软雅黑', size=None, color=None, bold=None):
    style.font.name = name
    if size is not None: style.font.size = Pt(size)
    if color is not None: style.font.color.rgb = color
    if bold is not None: style.font.bold = bool(bold)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name); rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name)

def get_style(names):
    for n in names:
        try: return doc.styles[n]
        except KeyError: continue
    return None

def add_field(run, code):
    r = run._r
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r.append(b)
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = code; r.append(i)
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r.append(s)
    t = OxmlElement('w:t'); t.text = '1'; r.append(t)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r.append(e)

def insert_para_before(ref_el):
    from docx.text.paragraph import Paragraph
    ref = ref_el._p if hasattr(ref_el, '_p') else ref_el
    p = OxmlElement('w:p'); ref.addprevious(p); return Paragraph(p, doc)

def set_cell_shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

# ---------- 1) 全局 Normal 字体 ----------
normal = doc.styles['Normal']
style_set_font(normal, size=10.5, color=DARK)

# ---------- 2) 标题样式 ----------
h1 = get_style(['Heading 1', '标题 1'])
h2 = get_style(['Heading 2', '标题 2'])
h3 = get_style(['Heading 3', '标题 3'])
has_h = h1 is not None
if has_h:
    style_set_font(h1, size=16, color=GREEN, bold=True)
    style_set_font(h2, size=13, color=DARK, bold=True)
    style_set_font(h3, size=11.5, color=DARK, bold=True)
    h1.paragraph_format.space_before = Pt(12); h1.paragraph_format.space_after = Pt(6); h1.paragraph_format.page_break_before = True
    h2.paragraph_format.space_before = Pt(8);  h2.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.space_before = Pt(6);  h3.paragraph_format.space_after = Pt(2)

# ---------- 3) 封面美化 ----------
title_p = next((p for p in doc.paragraphs if p.text.strip() == '综合场站大宗骨料智慧管控与检测系统'), None)
sub_p   = next((p for p in doc.paragraphs if p.text.strip() == '产品介绍'), None)
foot_p  = next((p for p in doc.paragraphs if '山东高速工程装备科技有限公司' in p.text), None)

if title_p:
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in title_p.runs: set_run_font(r, size=26, color=DARK, bold=True)
    pPr = title_p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '24'); bottom.set(qn('w:space'), '6'); bottom.set(qn('w:color'), '22C55E')
    pBdr.append(bottom); pPr.append(pBdr)
if sub_p:
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub_p.runs: set_run_font(r, size=16, color=GREEN, bold=True)
if foot_p:
    foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in foot_p.runs: set_run_font(r, size=10.5, color=GRAY)
    run = foot_p.add_run(); br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); run._r.append(br)

# ---------- 4) 自动目录 替换手动目录 ----------
paras = doc.paragraphs
idx_toc = idx_h1 = None
for i, p in enumerate(paras):
    t = p.text.strip()
    if idx_toc is None and (t == '目  录' or t == '目录'):
        idx_toc = i
    if idx_toc is not None and idx_h1 is None and re.match(r'^[一二三四]、[^—]+$', t):
        idx_h1 = i; break
if idx_toc is not None and idx_h1 is not None:
    for i in range(idx_h1 - 1, idx_toc - 1, -1):
        el = doc.paragraphs[i]._element
        el.getparent().remove(el)
    ref = doc.paragraphs[idx_toc]._element
    toc_p = insert_para_before(ref)
    run = toc_p.add_run()
    add_field(run, 'TOC \\o "1-3" \\h \\z \\u')
    title_p2 = insert_para_before(toc_p)
    run = title_p2.add_run('目  录')
    set_run_font(run, size=16, color=GREEN, bold=True)

# ---------- 5) 标题层级 + 图注 + 图片居中 ----------
for p in doc.paragraphs:
    t = p.text.strip()
    if not t: continue
    if re.match(r'^[一二三四五六]、', t) and '—' not in t:
        if has_h: p.style = h1
        continue
    if re.match(r'^\d+\.\d+', t):
        if has_h: p.style = h2
        continue
    if t in ('面向行业', '产品形态：三位一体智能化装备系统', '建设目标', '核心特性', '核心功能'):
        if has_h: p.style = h3
        continue
    if t.startswith('图：'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: set_run_font(r, size=9, color=GRAY)
        continue
for p in doc.paragraphs:
    if p._element.find('.//' + qn('w:drawing')) is not None:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- 6) 表格美化 ----------
if doc.tables:
    tbl = doc.tables[0]
    for cell in tbl.rows[0].cells:
        set_cell_shade(cell, '22C55E')
        for para in cell.paragraphs:
            for r in para.runs:
                r.font.color.rgb = WHITE; r.font.bold = True
                set_run_font(r, size=10)
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'C8CDD4'); borders.append(e)
    tbl_pr.append(borders)
    for ri, row in enumerate(tbl.rows[1:], start=1):
        if ri % 2 == 0:
            for cell in row.cells:
                set_cell_shade(cell, 'F2F7F2')

# ---------- 7) 页眉页脚 ----------
section = doc.sections[0]
hp = section.header.paragraphs[0]; hp.text = ''
run = hp.add_run('山东高速工程装备科技有限公司'); set_run_font(run, size=9, color=GRAY)
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp = section.footer.paragraphs[0]; fp.text = ''; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('第 '); add_field(fp.add_run(), 'PAGE'); fp.add_run(' 页 / 共 '); add_field(fp.add_run(), 'NUMPAGES'); fp.add_run(' 页')
for r in fp.runs: set_run_font(r, size=9, color=GRAY)

# ---------- 8) 打开时自动更新域 ----------
settings = doc.settings.element
uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); settings.append(uf)

doc.save(DST)
print("✅ 已生成:", DST)
print("   Heading样式:", "可用" if has_h else "缺失(已降级)")
print("   大小: %.1f KB" % (len(open(DST, 'rb').read()) / 1024))
