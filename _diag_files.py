from docx import Document
from pptx import Presentation
from pptx.util import Emu

def diag_docx(path):
    print("#" * 70)
    print("DOCX:", path)
    doc = Document(path)
    print(f"段落数={len(doc.paragraphs)} 图片数={len(doc.inline_shapes)} 表格数={len(doc.tables)}")
    print("--- 正文大纲(标题/正文) ---")
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        st = p.style.name if p.style else "?"
        print(f"[{st}] {txt[:60]}")
    print("--- 表格行数 ---")
    for i, t in enumerate(doc.tables, 1):
        print(f"  表{i}: {len(t.rows)}行 x {len(t.columns)}列")

def diag_pptx(path):
    print("#" * 70)
    print("PPTX:", path)
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        print(f"=== Slide {i} ===")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                # 合并多行
                t = " / ".join(l.strip() for l in sh.text_frame.text.splitlines() if l.strip())
                if t:
                    print(f"  · {t[:70]}")

import sys
if sys.argv[1].endswith(".docx"):
    diag_docx(sys.argv[1])
elif sys.argv[1].endswith(".pptx"):
    diag_pptx(sys.argv[1])
