from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

DST = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍_整理版.docx'
SRC = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.docx'
GREEN = RGBColor(0x22, 0xC5, 0x5E)

def set_run_font(run, name='微软雅黑', size=None, color=None, bold=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name); rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name)
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.font.bold = bool(bold)

def add_field(run, code):
    r = run._r
    for tag, val in (('begin', None), ('separate', '1'), ('end', None)):
        fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tag); r.append(fc)
        if val is not None:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code; r.append(it)
            sp = OxmlElement('w:fldChar'); sp.set(qn('w:fldCharType'), 'separate'); r.append(sp)
            t = OxmlElement('w:t'); t.text = val; r.append(t)

dst = Document(DST)

# 1) 修正图2：删除被放大的图(宽>5in)段落，从原文档取原图重新以 2.22in 嵌入
for p in list(dst.paragraphs):
    if p._element.find('.//' + qn('w:drawing')) is not None:
        ext = p._element.find('.//' + qn('wp:extent'))
        if ext is not None and int(ext.get('cx')) / 914400 > 5:
            p._element.getparent().remove(p._element)
            print('已删除放大版图2段落')

src = Document(SRC)
blip = src.inline_shapes[1]._inline.find('.//' + qn('a:blip'))
embed = blip.get(qn('r:embed'))
blob = src.part.related_parts[embed].blob
cap = next(p for p in dst.paragraphs if p.text.strip().startswith('图：智能盘点机器人'))
new_p = OxmlElement('w:p'); cap._p.addprevious(new_p)
np_ = Paragraph(new_p, dst); run = np_.add_run()
run.add_picture(BytesIO(blob), width=Inches(2.22))
np_.alignment = WD_ALIGN_PARAGRAPH.CENTER
print('已重插图2 (2.22in)')

# 2) 重建目录：在第一个 H1 前插入 目录标题 + TOC 域
h1 = next(p for p in dst.paragraphs if p.style.name == 'Heading 1' and p.text.strip().startswith('一、'))
ref = h1._p
toc_p = OxmlElement('w:p'); ref.addprevious(toc_p)
toc_para = Paragraph(toc_p, dst); trun = toc_para.add_run()
add_field(trun, 'TOC \\o "1-3" \\h \\z \\u')
title_p = OxmlElement('w:p'); toc_p.addprevious(title_p)
title_para = Paragraph(title_p, dst); t2 = title_para.add_run('目  录')
set_run_font(t2, size=16, color=GREEN, bold=True)
print('已重建目录(标题+TOC域)')

# 3) 自动更新域
settings = dst.settings.element
uf = settings.find(qn('w:updateFields'))
if uf is None:
    uf = OxmlElement('w:updateFields'); settings.append(uf)
uf.set(qn('w:val'), 'true')

dst.save(DST)
print('✅ 保存完成')
