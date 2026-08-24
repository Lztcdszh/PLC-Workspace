from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

SRC = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.docx'
DST = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍_整理版.docx'

DARK = RGBColor(0x1F, 0x2A, 0x37)

def set_run_font(run, name='微软雅黑', size=None, color=None, bold=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name); rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name)
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.font.bold = bool(bold)

src = Document(SRC)
# 通过 r:embed 关系ID 提取图片字节
blip = src.inline_shapes[1]._inline.find('.//' + qn('a:blip'))
embed = blip.get(qn('r:embed'))
blob = src.part.related_parts[embed].blob
print('提取到原图2, bytes=', len(blob))

dst = Document(DST)

# ---- 1) 补回丢失的机器人配图：插到"图：智能盘点机器人"图注之前 ----
cap = next(p for p in dst.paragraphs if p.text.strip().startswith('图：智能盘点机器人'))
new_p = OxmlElement('w:p')
cap._p.addprevious(new_p)
np_ = Paragraph(new_p, dst)
run = np_.add_run()
run.add_picture(BytesIO(blob))   # 保持原始尺寸 2.22x1.89
np_.alignment = WD_ALIGN_PARAGRAPH.CENTER
print('已插入图2，位于图注前')

# ---- 2) 恢复"核心功能"小标题文字 ----
fixed = False
for p in dst.paragraphs:
    if p.style.name == 'Heading 3' and not p.text.strip():
        r = p.add_run('核心功能')
        set_run_font(r, size=11.5, color=DARK, bold=True)
        fixed = True
        break
print('恢复核心功能文字:', fixed)

dst.save(DST)
print('✅ 修补完成并保存')
