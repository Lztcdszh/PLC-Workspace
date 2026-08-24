# -*- coding: utf-8 -*-
"""v3: 工业极简 + 产品发布风，渐变/阴影/图标层，覆盖生成 8 页。"""
import os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ASSET = r"C:\Users\Administrator\source\repos\_slide_assets"
OUT = r"C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.pptx"

WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x22, 0xC5, 0x5E)
DGREEN = RGBColor(0x06, 0x4E, 0x3B)
DARK   = RGBColor(0x1A, 0x22, 0x33)
GRAY   = RGBColor(0x6B, 0x72, 0x80)
DARKBG = RGBColor(0x0B, 0x14, 0x24)
DARKBG2= RGBColor(0x0E, 0x2A, 0x3F)
FONT   = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = Inches(13.333); H = Inches(7.5)
MX = 0.7


def _ea(run, font=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def hx(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def set_gradient(shape, h1, h2, angle=90):
    sp = shape.element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:noFill", "a:grpFill"):
        e = sp.find(qn(tag))
        if e is not None:
            sp.remove(e)
    gf = etree.SubElement(sp, qn("a:gradFill"))
    gsLst = etree.SubElement(gf, qn("a:gsLst"))
    for pos, c in ((0, hx(h1)), (100000, hx(h2))):
        gs = etree.SubElement(gsLst, qn("a:gs"))
        gs.set("pos", str(pos))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", "%02X%02X%02X" % c)
    lin = etree.SubElement(gf, qn("a:lin"))
    lin.set("ang", str(int(angle * 60000)))
    lin.set("scaled", "1")


def add_shadow(shape, blur=95000, dist=38000, dir=5400000, alpha=0.16):
    sp = shape.element.spPr
    el = sp.find(qn("a:effectLst"))
    if el is not None:
        sp.remove(el)
    eff = etree.SubElement(sp, qn("a:effectLst"))
    shdw = etree.SubElement(eff, qn("a:outerShdw"))
    shdw.set("blurRad", str(blur))
    shdw.set("dist", str(dist))
    shdw.set("dir", str(dir))
    shdw.set("rotWithShape", "0")
    clr = etree.SubElement(shdw, qn("a:srgbClr"))
    clr.set("val", "0B1220")
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def rect(slide, x, y, w, h, fill=None, line=None, line_w=None, radius=None, shadow=False):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    if radius:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if shadow:
        add_shadow(sp)
    return sp


def oval(slide, x, y, w, h, fill=None, line=None, line_w=None, shadow=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    if shadow:
        add_shadow(sp)
    return sp


def runs(slide, x, y, w, h, spec, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.12, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for txt, sz, col, bold in spec:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.name = FONT; r.font.color.rgb = col
        _ea(r)
    return tb


def lines(slide, x, y, w, h, text, size, color=DARK, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = FONT; r.font.color.rgb = color
        _ea(r)
    return tb


def img(slide, path, x, y, w=None, h=None, shadow=False):
    p = slide.shapes.add_picture(path, x, y, width=w, height=h)
    if shadow:
        add_shadow(p)
    return p


def deco_dots(slide, x, y, scale=1.0):
    oval(slide, x, y, Inches(0.18*scale), Inches(0.18*scale), fill=GREEN)
    oval(slide, x+Inches(0.28*scale), y+Inches(0.05*scale), Inches(0.11*scale), Inches(0.11*scale), fill=GREEN)
    oval(slide, x+Inches(0.45*scale), y+Inches(0.10*scale), Inches(0.07*scale), Inches(0.07*scale), fill=GREEN)


def corner_deco(slide, x=Inches(11.2), y=Inches(0.35)):
    oval(slide, x, y, Inches(1.4), Inches(1.4), line=GREEN, line_w=Pt(1.3))
    deco_dots(slide, x+Inches(0.45), y+Inches(0.35), scale=0.9)


def bg(slide, h1, h2, angle=90):
    r = rect(slide, 0, 0, W, H, fill=WHITE)
    set_gradient(r, h1, h2, angle)
    return r


def page_header(slide, no, label):
    bar = rect(slide, 0, 0, W, Inches(0.14), fill=GREEN)
    set_gradient(bar, "#22C55E", "#10B981", 90)
    runs(slide, MX, Inches(0.42), Inches(8), Inches(0.5),
         [(no, 24, GREEN, True), ("   " + label, 22, DARK, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    corner_deco(slide)


def img_card(slide, path, x, y, w):
    card = rect(slide, x-Inches(0.18), y-Inches(0.18), w+Inches(0.36), w*0.74+Inches(0.36),
                fill=WHITE, radius=0.06, shadow=True)
    img(slide, path, x, y, w=w)
    return card


def icon_block(slide, x, y, sym="✓"):
    rect(slide, x, y, Inches(0.5), Inches(0.5), fill=GREEN, radius=0.28, shadow=True)
    runs(slide, x, y+Inches(0.02), Inches(0.5), Inches(0.46),
         [(sym, 22, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# =========================================================
# Slide 1 — 封面
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
rect(s, Inches(0.7), Inches(2.35), Inches(0.14), Inches(2.6), fill=GREEN, shadow=True)
lines(s, Inches(1.1), Inches(2.32), Inches(10.5), Inches(0.95), "综合场站大宗骨料", 50, DARK, True)
lines(s, Inches(1.1), Inches(3.3), Inches(10.5), Inches(0.95), "智慧管控与检测系统", 50, DARK, True)
runs(s, Inches(1.14), Inches(4.6), Inches(6), Inches(0.6), [("产品介绍", 32, GREEN, True)])
rect(s, Inches(1.14), Inches(5.34), Inches(2.4), Inches(0.05), fill=GREEN)
lines(s, Inches(1.14), Inches(6.5), Inches(10), Inches(0.5),
      "山东高速工程装备科技有限公司   |   2026", 18, GRAY)
oval(s, Inches(10.4), Inches(0.7), Inches(2.6), Inches(2.6), line=GREEN, line_w=Pt(1.5))
deco_dots(s, Inches(11.1), Inches(1.2), scale=1.2)
oval(s, Inches(0.7), Inches(5.0), Inches(0.55), Inches(0.55), fill=GREEN, shadow=True)
runs(s, Inches(0.55), Inches(5.06), Inches(0.6), Inches(0.4), [("智", 16, WHITE, True)], align=PP_ALIGN.CENTER)

# =========================================================
# Slide 2 — 产品定位
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
page_header(s, "01", "产品定位")
lines(s, Inches(0.7), Inches(1.5), Inches(7.2), Inches(1.5),
      "「智慧骨料管家」\n骨料管理的“数字大脑”", 38, DARK, True, line_spacing=1.18)
points = [
    ("自动盘点", "库存自动采集，告别人工估算"),
    ("质量在线检测", "全过程质量数据实时可溯"),
    ("数据闭环", "采集—分析—管控全链路打通"),
    ("智能调度管控", "料仓状态可视，调度有据可依"),
]
py = 3.55
for t, d in points:
    icon_block(s, Inches(0.7), Inches(py), "✓")
    runs(s, Inches(1.4), Inches(py-0.04), Inches(6), Inches(0.45), [(t, 22, DARK, True)])
    runs(s, Inches(1.4), Inches(py+0.38), Inches(6), Inches(0.4), [(d, 15, GRAY, False)])
    py += 0.82
lines(s, Inches(0.7), Inches(6.95), Inches(7), Inches(0.5),
      "面向 混凝土搅拌站 · 建材综合场站 · 砂石骨料加工", 16, GRAY)
img_card(s, os.path.join(ASSET, "s06_p10_236x236.png"), Inches(8.2), Inches(2.1), Inches(4.4))

# =========================================================
# Slide 3 — 解决方案总览
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
page_header(s, "02", "解决方案")
img_card(s, os.path.join(ASSET, "s07_p12_613x453.png"), Inches(0.7), Inches(1.6), Inches(6.0))
lines(s, Inches(7.2), Inches(1.6), Inches(5.4), Inches(1.4),
      "三位一体\n智能化装备系统", 38, DARK, True, line_spacing=1.18)
three = [
    ("硬件感知", "导轨机器人 + 激光雷达，料堆三维采集"),
    ("算法决策", "点云处理 + AI，体积与质量精准测算"),
    ("平台管控", "可视化看板 + 报表，全流程数字化"),
]
py = 3.6
for t, d in three:
    rect(s, Inches(7.2), Inches(py), Inches(0.16), Inches(0.85), fill=GREEN, shadow=True)
    runs(s, Inches(7.55), Inches(py-0.02), Inches(5), Inches(0.4), [(t, 23, DARK, True)])
    runs(s, Inches(7.55), Inches(py+0.38), Inches(5), Inches(0.5), [(d, 15, GRAY, False)])
    py += 1.0
lines(s, Inches(7.2), Inches(6.75), Inches(5.4), Inches(0.5),
      "融合 激光雷达 · 三维点云 · 工业物联网 · AI", 15, GREEN)

# =========================================================
# Slide 4 — 核心硬件
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
page_header(s, "03", "核心硬件")
lines(s, Inches(0.7), Inches(1.0), Inches(9), Inches(0.7), "智能盘点机器人", 42, DARK, True)
runs(s, Inches(0.72), Inches(1.78), Inches(9), Inches(0.4),
     [("硬件基石 · 自动巡航，精准采集料堆体积", 19, GRAY, False)])
img_card(s, os.path.join(ASSET, "s06_p10_236x236.png"), Inches(0.95), Inches(2.5), Inches(3.5))
cards = [
    ("自动化作业", "7×24 小时无人值守，按预设路径自动采集"),
    ("数据全覆盖", "180° 激光雷达扫描，无死角获取三维数据"),
    ("高适应性", "滑触线 / 电池双供电，适应多种料场工况"),
]
cy = 2.45
for t, d in cards:
    rect(s, Inches(4.9), Inches(cy), Inches(7.7), Inches(1.0), fill=WHITE, radius=0.08, shadow=True)
    rect(s, Inches(4.9), Inches(cy), Inches(0.14), Inches(1.0), fill=GREEN)
    runs(s, Inches(5.25), Inches(cy+0.14), Inches(7), Inches(0.4), [(t, 21, DARK, True)])
    runs(s, Inches(5.25), Inches(cy+0.55), Inches(7.2), Inches(0.4), [(d, 14, GRAY, False)])
    cy += 1.13
traj = [("s06_p14_222x138.png", "U 型巡检轨迹"),
        ("s06_p15_222x138.png", "一字型巡检轨迹"),
        ("s06_p16_222x138.png", "定制化巡检轨迹")]
tx = 0.7
for fn, lb in traj:
    rect(s, Inches(tx), Inches(6.0), Inches(3.7), Inches(1.0), fill=WHITE, radius=0.06, shadow=True)
    img(s, os.path.join(ASSET, fn), Inches(tx+0.2), Inches(6.18), w=Inches(1.5))
    runs(s, Inches(tx+1.85), Inches(6.4), Inches(1.7), Inches(0.5), [(lb, 14, DARK, True)])
    tx += 4.06

# =========================================================
# Slide 5 — 核心软件
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
page_header(s, "04", "核心软件")
lines(s, Inches(0.7), Inches(1.0), Inches(9), Inches(0.7), "智慧管理平台", 42, DARK, True)
runs(s, Inches(0.72), Inches(1.78), Inches(9), Inches(0.4),
     [("软件中枢 · 打通生产与管理数据壁垒", 19, GRAY, False)])
img_card(s, os.path.join(ASSET, "s07_p12_613x453.png"), Inches(0.8), Inches(2.4), Inches(6.0))
funcs = [
    ("实时可视化看板", "三维模型 + 动态图表，直观展示各料仓库存与状态"),
    ("自动化报表生成", "日报 / 月报自动生成，支持多维统计与历史追溯"),
    ("开放 API 接口", "无缝对接 ERP / MES，实现企业数据全面互通"),
]
fy = 2.45
for t, d in funcs:
    rect(s, Inches(7.2), Inches(fy), Inches(5.4), Inches(1.25), fill=WHITE, radius=0.08, shadow=True)
    rect(s, Inches(7.2), Inches(fy), Inches(0.14), Inches(1.25), fill=GREEN)
    runs(s, Inches(7.55), Inches(fy+0.2), Inches(5), Inches(0.4), [(t, 22, DARK, True)])
    runs(s, Inches(7.55), Inches(fy+0.66), Inches(5), Inches(0.55), [(d, 14, GRAY, False)])
    fy += 1.4

# =========================================================
# Slide 6 — 核心技术优势 (深色)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#0B1424", "#0E2A3F", 120)
oval(s, Inches(10.4), Inches(0.6), Inches(2.6), Inches(2.6), line=GREEN, line_w=Pt(1.5))
deco_dots(s, Inches(11.1), Inches(1.1), scale=1.1)
oval(s, Inches(0.5), Inches(5.0), Inches(2.0), Inches(2.0), line=GREEN, line_w=Pt(1.5))
runs(s, Inches(1.0), Inches(2.55), Inches(11.33), Inches(1.2),
     [("全栈式「", 46, WHITE, True),
      ("感知", 46, GREEN, True), ("—", 46, WHITE, True),
      ("分析", 46, GREEN, True), ("—", 46, WHITE, True),
      ("管控", 46, GREEN, True), ("」一体化架构", 46, WHITE, True)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
runs(s, Inches(1.0), Inches(4.1), Inches(11.33), Inches(0.6),
     [("全链路自主可控 · 软硬一体 · 数据闭环", 22, RGBColor(0x9C,0xA3,0xAF), False)],
     align=PP_ALIGN.CENTER)

# =========================================================
# Slide 7 — 技术优势对比 (表格)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#F5F8FB", "#FFFFFF", 135)
page_header(s, "05", "核心优势")
lines(s, Inches(0.7), Inches(1.0), Inches(9), Inches(0.7), "技术优势对比", 36, DARK, True)
rows = [
    ["对比维度", "传统人工盘点", "智慧骨料管家系统"],
    ["盘点精度", "误差 > 15%（人工估算）", "误差 ≤ 5%（激光精准测量）"],
    ["盘点效率", "人工盘点，效率低下", "全自动盘点（定时 / 定点）"],
    ["人力成本", "需 3–5 名专职盘点员", "无人值守，极大节省人力"],
    ["数据实时性", "数据滞后，呈静态状态", "实时更新，库存动态监测"],
    ["安全风险 / 数据价值", "高风险（登高作业）/ 单一孤立", "无风险（非接触式）/ 多维可追溯"],
]
rect(s, Inches(0.6), Inches(1.85), Inches(12.13), Inches(5.05), fill=WHITE, radius=0.05, shadow=True)
nrows, ncols = len(rows), 3
tbl_shape = s.shapes.add_table(nrows, ncols, Inches(0.85), Inches(2.05),
                               Inches(11.63), Inches(4.7))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.3)
tbl.columns[1].width = Inches(4.165)
tbl.columns[2].width = Inches(4.165)
for ri, row in enumerate(rows):
    tbl.rows[ri].height = Inches(0.78)
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.18); cell.margin_right = Inches(0.12)
        cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.name = FONT; _ea(r)
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = DGREEN
            r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = WHITE
        else:
            if ci == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF1,0xF5,0xF9)
                r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = DARK
            elif ci == 1:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF1,0xF5,0xF9)
                r.font.size = Pt(14); r.font.color.rgb = GRAY
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xEC,0xFD,0xF5)
                r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = DGREEN

# =========================================================
# Slide 8 — 致谢 (深色)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, "#0B1424", "#0E2A3F", 120)
rect(s, Inches(5.0), Inches(2.35), Inches(3.33), Inches(0.06), fill=GREEN)
runs(s, Inches(1.0), Inches(2.6), Inches(11.33), Inches(1.2),
     [("THANKS", 80, GREEN, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
runs(s, Inches(1.0), Inches(3.95), Inches(11.33), Inches(0.6),
     [("感 谢 聆 听", 28, WHITE, False)], align=PP_ALIGN.CENTER)
runs(s, Inches(1.0), Inches(4.75), Inches(11.33), Inches(0.5),
     [("山东高速工程装备科技有限公司", 18, RGBColor(0x9C,0xA3,0xAF), False)], align=PP_ALIGN.CENTER)
deco_dots(s, Inches(11.0), Inches(5.5), scale=1.0)
oval(s, Inches(0.7), Inches(0.7), Inches(1.4), Inches(1.4), line=GREEN, line_w=Pt(1.5))

prs.save(OUT)
print("Saved:", OUT, "Slides:", len(prs.slides._sldIdLst))
