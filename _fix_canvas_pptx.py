"""修复画布版PPTX：将导出的PNG以正确的全屏尺寸重新嵌入"""
from pptx import Presentation
from pptx.util import Inches
import os

SLIDE_W = 13.333
SLIDE_H = 7.500
EXPORT_DIR = r"C:\Users\Administrator\source\repos\_slide_exports"
OUTPUT = r"C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍_画布版.pptx"

# 按文件名排序: 3_2.png ~ 3_9.png → 8页
png_files = sorted(f for f in os.listdir(EXPORT_DIR) if f.lower().endswith(".png"))
print(f"共 {len(png_files)} 张截图: {png_files}")

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
blank_layout = prs.slide_layouts[6]  # blank

for i, png_name in enumerate(png_files, 1):
    slide = prs.slides.add_slide(blank_layout)
    path = os.path.join(EXPORT_DIR, png_name)
    # 关键：明确指定 left/top/width/height 全屏铺满
    pic = slide.shapes.add_picture(path, Inches(0), Inches(0), width=Inches(SLIDE_W), height=Inches(SLIDE_H))
    print(f"  Slide {i}: {png_name} → {pic.width/914400:.3f}x{pic.height/914400:.3f} in")

prs.save(OUTPUT)
print(f"\n✅ 已保存: {OUTPUT}")
print(f"   大小: {os.path.getsize(OUTPUT)/1024/1024:.1f} MB")
