from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

SRC = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.docx'
DST = r'C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍_整理版.docx'

def info(name, path):
    d = Document(path)
    print(f'\n== {name}  图片{len(d.inline_shapes)}张 ==')
    for idx, shp in enumerate(d.inline_shapes, 1):
        w = Emu(shp.width).inches; h = Emu(shp.height).inches
        cap = ''
        for j, p in enumerate(d.paragraphs):
            if shp._inline.getparent().getparent() is p._p:
                cap = d.paragraphs[j+1].text.strip() if j+1 < len(d.paragraphs) else ''
                break
        print(f'  图{idx}: {w:.2f}x{h:.2f}in  图注={cap!r}')

info('原文档', SRC)
info('整理版', DST)

print('\n== 核心功能 小标题 ==')
d = Document(SRC)
for p in d.paragraphs:
    if p.text.strip() == '核心功能':
        print('  原文档:', repr(p.text), 'runs=', [r.text for r in p.runs])
d2 = Document(DST)
for i, p in enumerate(d2.paragraphs):
    if p.style.name == 'Heading 3' and not p.text.strip():
        prev = d2.paragraphs[i-1].text.strip()[:20] if i > 0 else ''
        nxt = d2.paragraphs[i+1].text.strip()[:20] if i+1 < len(d2.paragraphs) else ''
        print(f'  整理版段{i}: 空H3 | 上={prev!r} 下={nxt!r}')
