# -*- coding: utf-8 -*-
"""重建 8 页产品介绍 PPT — 工业极简风 + 品牌绿 (#22C55E)，发布会演讲场景。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ASSET = r"C:\Users\Administrator\source\repos\_slide_assets"
OUT = r"C:\Users\Administrator\source\repos\综合场站大宗骨料智慧管控与检测系统_产品介绍.pptx"

# ---- design tokens ----
BG        = RGBColor(0xF8, 0xFA, 0xFC)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREEN     = RGBColor(0x22, 0xC5, 0x5E)
GREEND    = RGBColor(0x10, 0x98, 0x81)
DGREEN    = RGBColor(0x06, 0x4E, 0x3B)
DARK      = RGBColor(0x1F, 0x29, 0x37)
GRAY      = RGBColor(0x6B, 0x72, 0x80)
LIGHT     = RGBColor(0xE5, 0xE7, 0xEB)
DARKBG    = RGBColor(0x0F, 0x17, 0x2A)
ROWGRY    = RGBColor(0xF1, 0xF5, 0xF9)
ROWGRN    = RGBColor(0xEC, 0xFD, 0xF5)
FONT      = "微软雅黑"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = Inches(13.333); H = Inches(7.5)
MX = 0.7


def _ea(run, font):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=None, radius=None):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    if radius:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def oval(slide, x, y, w, h, fill=None, line=None, line_w=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    return sp


def runs(slide, x, y, w, h, spec, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.12, wrap=True):
    """spec: list of (text, size, color, bold) -> runs (同一段)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for i, (txt, sz, col, bold) in enumerate(spec):
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.name = FONT; r.font.color.rgb = col
        _ea(r, FONT)
    return tb


def lines(slide, x, y, w, h, text, size, color=DARK, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    """多行纯文本（\\n 分隔）"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = FONT; r.font.color.rgb = color
        _ea(r, FONT)
    return tb


def img(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def bg(slide, color):
    return rect(slide, 0, 0, W, H, fill=color)


def deco_dots(slide, x, y, scale=1.0):
    oval(slide, x, y, Inches(0.18*scale), Inches(0.18*scale), fill=GREEN)
    oval(slide, x+Inches(0.28*scale), y+Inches(0.05*scale), Inches(0.11*scale), Inches(0.11*scale), fill=GREEN)
    oval(slide, x+Inches(0.45*scale), y+Inches(0.10*scale), Inches(0.07*scale), Inches(0.07*scale), fill=GREEN)


def tag(slide, no, label, y=0.55):
    """页面小标签：序号 + 文字"""
    runs(slide, MX, Inches(y), Inches(6), Inches(0.5),
         [(no, 22, GREEN, True), ("   " + label, 22, DARK, True)],
         anchor=MSO_ANCHOR.MIDDLE)


# =========================================================
# Slide 1 — 封面
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
rect(s, Inches(0.7), Inches(2.35), Inches(0.13), Inches(2.5), fill=GREEN)
lines(s, Inches(1.08), Inches(2.35), Inches(10.5), Inches(0.95),
      "综合场站大宗骨料", 48, DARK, True)
lines(s, Inches(1.08), Inches(3.32), Inches(10.5), Inches(0.95),
      "智慧管控与检测系统", 48, DARK, True)
runs(s, Inches(1.12), Inches(4.55), Inches(6), Inches(0.6),
     [("产品介绍", 30, GREEN, True)])
rect(s, Inches(1.12), Inches(5.28), Inches(2.3), Inches(0.04), fill=RGBColor(0x22,0xC5,0x5E))
lines(s, Inches(1.12), Inches(6.5), Inches(10), Inches(0.5),
      "山东高速工程装备科技有限公司   |   2026", 18, GRAY)
# 右上角装饰
oval(s, Inches(10.3), Inches(0.7), Inches(2.6), Inches(2.6), line=RGBColor(0x22,0xC5,0x5E), line_w=Pt(1.5))
deco_dots(s, Inches(11.9), Inches(1.2), scale=1.2)
oval(s, Inches(0.7), Inches(5.0), Inches(0.5), Inches(0.5), fill=RGBColor(0x22,0xC5,0x5E))
runs(s, Inches(0.55), Inches(5.05), Inches(1), Inches(0.4), [("智", 16, WHITE, True)])

# =========================================================
# Slide 2 — 产品定位 (左文右图)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
tag(s, "01", "产品定位")
lines(s, Inches(0.7), Inches(1.45), Inches(7.2), Inches(1.5),
      "「智慧骨料管家」\n骨料管理的“数字大脑”", 36, DARK, True, line_spacing=1.18)
points = [
    ("自动盘点", "库存自动采集，告别人工估算"),
    ("质量在线检测", "全过程质量数据实时可溯"),
    ("数据闭环", "采集—分析—管控全链路打通"),
    ("智能调度管控", "料仓状态可视，调度有据可依"),
]
py = 3.55
for i, (t, d) in enumerate(points):
    rect(s, Inches(0.7), Inches(py), Inches(0.32), Inches(0.32), fill=GREEN, radius=0.5)
    runs(s, Inches(1.18), Inches(py-0.05), Inches(6), Inches(0.45),
         [(t, 21, DARK, True)])
    runs(s, Inches(1.18), Inches(py+0.34), Inches(6), Inches(0.4),
         [(d, 15, GRAY, False)])
    py += 0.82
lines(s, Inches(0.7), Inches(6.75), Inches(7), Inches(0.5),
      "面向 混凝土搅拌站 · 建材综合场站 · 砂石骨料加工", 16, GRAY)
# 右侧机器人主视觉
img(s, os.path.join(ASSET, "s06_p10_236x236.png"), Inches(8.0), Inches(1.9), w=Inches(4.6))
rect(s, Inches(7.7), Inches(1.6), Inches(5.0), Inches(5.0), fill=WHITE, radius=0.06)
img(s, os.path.join(ASSET, "s06_p10_236x236.png"), Inches(8.05), Inches(2.0), w=Inches(4.3))

# =========================================================
# Slide 3 — 解决方案总览 (左图右文)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
tag(s, "02", "解决方案")
rect(s, Inches(0.7), Inches(1.5), Inches(6.1), Inches(5.0), fill=WHITE, radius=0.06)
img(s, os.path.join(ASSET, "s07_p12_613x453.png"), Inches(1.0), Inches(2.4), w=Inches(5.5))
lines(s, Inches(7.2), Inches(1.5), Inches(5.4), Inches(1.4),
      "三位一体\n智能化装备系统", 36, DARK, True, line_spacing=1.18)
three = [
    ("硬件感知", "导轨机器人 + 激光雷达，料堆三维采集"),
    ("算法决策", "点云处理 + AI，体积与质量精准测算"),
    ("平台管控", "可视化看板 + 报表，全流程数字化"),
]
py = 3.5
for t, d in three:
    rect(s, Inches(7.2), Inches(py), Inches(0.14), Inches(0.85), fill=GREEN)
    runs(s, Inches(7.5), Inches(py-0.02), Inches(5), Inches(0.4), [(t, 22, DARK, True)])
    runs(s, Inches(7.5), Inches(py+0.38), Inches(5), Inches(0.5), [(d, 15, GRAY, False)])
    py += 1.0
lines(s, Inches(7.2), Inches(6.7), Inches(5.4), Inches(0.5),
      "融合 激光雷达 · 三维点云 · 工业物联网 · AI", 15, GREEN)

# =========================================================
# Slide 4 — 核心硬件：智能盘点机器人
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
tag(s, "03", "核心硬件")
lines(s, Inches(0.7), Inches(1.0), Inches(8), Inches(0.7),
      "智能盘点机器人", 40, DARK, True)
runs(s, Inches(0.72), Inches(1.75), Inches(8), Inches(0.4),
     [("硬件基石 · 自动巡航，精准采集料堆体积", 18, GRAY, False)])
# 左：机器人大图
rect(s, Inches(0.7), Inches(2.3), Inches(3.9), Inches(3.3), fill=WHITE, radius=0.06)
img(s, os.path.join(ASSET, "s06_p10_236x236.png"), Inches(0.95), Inches(2.7), w=Inches(3.4))
# 右：三特性卡片
cards = [
    ("自动化作业", "7×24 小时无人值守，按预设路径自动采集"),
    ("数据全覆盖", "180° 激光雷达扫描，无死角获取三维数据"),
    ("高适应性", "滑触线 / 电池双供电，适应多种料场工况"),
]
cy = 2.35
for t, d in cards:
    rect(s, Inches(4.9), Inches(cy), Inches(7.7), Inches(1.0), fill=WHITE, radius=0.08)
    rect(s, Inches(4.9), Inches(cy), Inches(0.12), Inches(1.0), fill=GREEN)
    runs(s, Inches(5.2), Inches(cy+0.14), Inches(7), Inches(0.4), [(t, 20, DARK, True)])
    runs(s, Inches(5.2), Inches(cy+0.55), Inches(7.2), Inches(0.4), [(d, 14, GRAY, False)])
    cy += 1.13
# 底部：三种轨迹
traj = [("s06_p14_222x138.png", "U 型巡检轨迹"),
        ("s06_p15_222x138.png", "一字型巡检轨迹"),
        ("s06_p16_222x138.png", "定制化巡检轨迹")]
tx = 0.7
for fn, lb in traj:
    rect(s, Inches(tx), Inches(5.95), Inches(3.7), Inches(1.05), fill=WHITE, radius=0.06)
    im = img(s, os.path.join(ASSET, fn), Inches(tx+0.2), Inches(6.1), w=Inches(1.5))
    runs(s, Inches(tx+1.85), Inches(6.35), Inches(1.7), Inches(0.5), [(lb, 14, DARK, True)])
    tx += 4.06

# =========================================================
# Slide 5 — 核心软件：智慧管理平台
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
tag(s, "04", "核心软件")
lines(s, Inches(0.7), Inches(1.0), Inches(8), Inches(0.7),
      "智慧管理平台", 40, DARK, True)
runs(s, Inches(0.72), Inches(1.75), Inches(8), Inches(0.4),
     [("软件中枢 · 打通生产与管理数据壁垒", 18, GRAY, False)])
rect(s, Inches(0.7), Inches(2.3), Inches(6.2), Inches(4.3), fill=WHITE, radius=0.06)
img(s, os.path.join(ASSET, "s07_p12_613x453.png"), Inches(0.95), Inches(2.9), w=Inches(5.7))
funcs = [
    ("实时可视化看板", "三维模型 + 动态图表，直观展示各料仓库存与状态"),
    ("自动化报表生成", "日报 / 月报自动生成，支持多维统计与历史追溯"),
    ("开放 API 接口", "无缝对接 ERP / MES，实现企业数据全面互通"),
]
fy = 2.4
for t, d in funcs:
    rect(s, Inches(7.2), Inches(fy), Inches(5.4), Inches(1.25), fill=WHITE, radius=0.08)
    rect(s, Inches(7.2), Inches(fy), Inches(0.12), Inches(1.25), fill=GREEN)
    runs(s, Inches(7.5), Inches(fy+0.18), Inches(5), Inches(0.4), [(t, 21, DARK, True)])
    runs(s, Inches(7.5), Inches(fy+0.62), Inches(5), Inches(0.55), [(d, 14, GRAY, False)])
    fy += 1.4

# =========================================================
# Slide 6 — 核心技术优势 (深色居中)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, DARKBG)
oval(s, Inches(10.6), Inches(0.6), Inches(2.6), Inches(2.6), line=RGBColor(0x22,0xC5,0x5E), line_w=Pt(1.5))
deco_dots(s, Inches(11.3), Inches(1.1), scale=1.1)
oval(s, Inches(0.6), Inches(5.2), Inches(2.0), Inches(2.0), line=RGBColor(0x22,0xC5,0x5E), line_w=Pt(1.5))
runs(s, Inches(1.0), Inches(2.55), Inches(11.3), Inches(1.2),
     [("全栈式「", 44, WHITE, True),
      ("感知", 44, GREEN, True), ("—", 44, WHITE, True),
      ("分析", 44, GREEN, True), ("—", 44, WHITE, True),
      ("管控", 44, GREEN, True), ("」一体化架构", 44, WHITE, True)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
runs(s, Inches(1.0), Inches(4.1), Inches(11.3), Inches(0.6),
     [("全链路自主可控 · 软硬一体 · 数据闭环", 22, RGBColor(0x9C,0xA3,0xAF), False)],
     align=PP_ALIGN.CENTER)

# =========================================================
# Slide 7 — 技术优势对比 (表格)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, BG)
tag(s, "05", "核心优势")
lines(s, Inches(0.7), Inches(1.0), Inches(9), Inches(0.7), "技术优势对比", 34, DARK, True)
rows = [
    ["对比维度", "传统人工盘点", "智慧骨料管家系统"],
    ["盘点精度", "误差 > 15%（人工估算）", "误差 ≤ 5%（激光精准测量）"],
    ["盘点效率", "人工盘点，效率低下", "全自动盘点（定时 / 定点）"],
    ["人力成本", "需 3–5 名专职盘点员", "无人值守，极大节省人力"],
    ["数据实时性", "数据滞后，呈静态状态", "实时更新，库存动态监测"],
    ["安全风险 / 数据价值", "高风险（登高作业）/ 单一孤立", "无风险（非接触式）/ 多维可追溯"],
]
ncols = 3; nrows = len(rows)
tbl_shape = s.shapes.add_table(nrows, ncols, Inches(MX), Inches(1.9),
                               Inches(11.93), Inches(4.95))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.4)
tbl.columns[1].width = Inches(4.265)
tbl.columns[2].width = Inches(4.265)
for ri, row in enumerate(rows):
    tbl.rows[ri].height = Inches(0.82)
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.18); cell.margin_right = Inches(0.12)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = val
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = DGREEN
            r.font.size = Pt(18); r.font.bold = True
            r.font.color.rgb = WHITE; r.font.name = FONT
        else:
            if ci == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = ROWGRY
                r.font.size = Pt(15); r.font.bold = True
                r.font.color.rgb = DARK
            elif ci == 1:
                cell.fill.solid(); cell.fill.fore_color.rgb = ROWGRY
                r.font.size = Pt(14); r.font.bold = False
                r.font.color.rgb = GRAY
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = ROWGRN
                r.font.size = Pt(14); r.font.bold = True
                r.font.color.rgb = DGREEN
            r.font.name = FONT
        _ea(r, FONT)

# =========================================================
# Slide 8 — 致谢 (深色收尾)
# =========================================================
s = prs.slides.add_slide(BLANK)
bg(s, DARKBG)
rect(s, Inches(5.0), Inches(2.35), Inches(3.33), Inches(0.06), fill=GREEN)
runs(s, Inches(1.0), Inches(2.6), Inches(11.33), Inches(1.2),
     [("THANKS", 76, GREEN, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
runs(s, Inches(1.0), Inches(3.95), Inches(11.33), Inches(0.6),
     [("感 谢 聆 听", 26, WHITE, False)], align=PP_ALIGN.CENTER)
runs(s, Inches(1.0), Inches(4.75), Inches(11.33), Inches(0.5),
     [("山东高速工程装备科技有限公司", 18, RGBColor(0x9C,0xA3,0xAF), False)],
     align=PP_ALIGN.CENTER)
deco_dots(s, Inches(11.2), Inches(5.6), scale=1.0)
oval(s, Inches(0.7), Inches(0.7), Inches(1.4), Inches(1.4), line=RGBColor(0x22,0xC5,0x5E), line_w=Pt(1.5))

prs.save(OUT)
print("Saved:", OUT)
print("Slides:", len(prs.slides._sldIdLst))
