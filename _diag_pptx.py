from pptx import Presentation
from pptx.util import Emu
import sys

def diag(path):
    print("=" * 60)
    print("FILE:", path)
    prs = Presentation(path)
    sw, sh = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    print(f"slide size: {sw:.3f} x {sh:.3f} in")
    for i, slide in enumerate(prs.slides, 1):
        pics = [s for s in slide.shapes if s.shape_type == 13]
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        print(f"  Slide {i}: shapes={len(slide.shapes)} pics={len(pics)} textboxes={len(texts)}")
        for p in pics:
            l, t = Emu(p.left).inches, Emu(p.top).inches
            w, h = Emu(p.width).inches, Emu(p.height).inches
            # 判断是否覆盖整页
            cover = (abs(l) < 0.05 and abs(t) < 0.05 and abs(w - sw) < 0.1 and abs(h - sh) < 0.1)
            print(f"      pic L={l:.2f} T={t:.2f} W={w:.2f} H={h:.2f}  fullcover={cover}")

for p in sys.argv[1:]:
    try:
        diag(p)
    except Exception as e:
        print("ERR", p, e)
